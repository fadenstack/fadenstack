"""Chat attachments get their text without Docling.

A PDF or Office file attached to a chat needed the Docling service: without
it the upload was refused, and with Docling down the file was stored with no
text. Now the local extractor reads it in both cases, and Docling is still
used when it answers.
"""

from __future__ import annotations

import io
import uuid
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from faden_gateway.db.models.gateway import ExtractionStatus
from faden_gateway.services.gateway import extraction
from faden_gateway.services.gateway.attachment_service import AttachmentError, ChatAttachmentService
from faden_gateway.settings import settings


def _pdf(pages: list[str]) -> bytes:
    """A small, valid PDF with one line of text per page."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # the page tree, filled in below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids: list[int] = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
        content_id = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id,
        )
        kids.append(len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids),
        len(kids),
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def _docx(*paragraphs: str) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "North"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Costs"
    ws.append(["Office", "Cost"])
    ws.append(["Lisbon", 1200])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pptx(*slides: str) -> bytes:
    prs = Presentation()
    for text in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = text
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class _Dao:
    def __init__(self) -> None:
        self.created: list[dict[str, Any]] = []

    async def list_attachments_for_session(self, *, session_id: uuid.UUID) -> list[Any]:
        return []

    async def attachment_stats(self, *, tenant_id: str, user_id: str | None = None) -> dict[str, int]:
        return {"total_count": 0, "total_bytes": 0}

    async def create_attachment(self, **kwargs: Any) -> dict[str, Any]:
        self.created.append(kwargs)
        return kwargs


class _Store:
    def __init__(self) -> None:
        self.keys: list[str] = []

    async def put_bytes(self, key: str, data: bytes) -> None:
        self.keys.append(key)


class _Docling:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.calls = 0

    async def convert(self, file_bytes: bytes, filename: str, *, max_pages: int | None = None) -> dict[str, Any]:
        self.calls += 1
        if self.fail:
            raise ConnectionError("docling is down")
        return {"content": "from docling", "metadata": {"page_count": 1}}


async def _upload(service: ChatAttachmentService, data: bytes, filename: str) -> dict[str, Any]:
    result = await service.upload(
        file_bytes=data,
        filename=filename,
        content_type="application/octet-stream",
        tenant_id="tenant-a",
        user_id="user-1",
        session_id=uuid.uuid4(),
    )
    return result["attachment"]


def _service(docling: _Docling | None = None) -> tuple[ChatAttachmentService, _Dao, _Store]:
    dao, store = _Dao(), _Store()
    service = ChatAttachmentService(dao=dao, file_store=store, docling_client=docling)  # type: ignore[arg-type]
    return service, dao, store


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("filename", "data", "expected"),
    [
        ("report.pdf", _pdf(["Lisbon closes in June"]), "Lisbon closes in June"),
        ("memo.docx", _docx("Lisbon closes in June"), "Lisbon closes in June"),
        ("costs.xlsx", _xlsx(), "Lisbon | 1200"),
        ("deck.pptx", _pptx("Lisbon closes in June"), "Lisbon closes in June"),
    ],
    ids=["pdf", "docx", "xlsx", "pptx"],
)
async def test_without_docling_documents_are_extracted_locally(filename: str, data: bytes, expected: str) -> None:
    service, _, _ = _service(docling=None)
    attachment = await _upload(service, data, filename)
    assert attachment["extraction_status"] == ExtractionStatus.COMPLETED
    assert expected in attachment["extracted_text"]


@pytest.mark.anyio
async def test_docling_failure_falls_back_to_the_local_extractor() -> None:
    docling = _Docling(fail=True)
    service, _, _ = _service(docling)
    attachment = await _upload(service, _docx("Lisbon closes in June"), "memo.docx")
    assert docling.calls == 1
    assert attachment["extraction_status"] == ExtractionStatus.COMPLETED
    assert "Lisbon closes in June" in attachment["extracted_text"]


@pytest.mark.anyio
async def test_docling_is_used_when_it_answers() -> None:
    docling = _Docling()
    service, _, _ = _service(docling)
    attachment = await _upload(service, _docx("ignored"), "memo.docx")
    assert attachment["extracted_text"] == "from docling"


@pytest.mark.anyio
async def test_local_extraction_reads_at_most_the_page_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "chat_attachment_max_pages", 2)
    service, _, _ = _service(docling=None)
    attachment = await _upload(service, _pdf(["Page one", "Page two", "Page three"]), "report.pdf")
    text = attachment["extracted_text"]
    assert "Page one" in text and "Page two" in text
    assert "Page three" not in text
    assert attachment["truncated"] is True
    assert attachment["page_count"] == 3
    assert "Only the first 2 of 3 pages" in text


@pytest.mark.anyio
async def test_a_pdf_without_text_is_marked_failed_not_empty() -> None:
    service, _, _ = _service(docling=None)
    attachment = await _upload(service, _pdf([""]), "scan.pdf")
    assert attachment["extraction_status"] == ExtractionStatus.FAILED
    assert attachment["extracted_text"] is None


@pytest.mark.anyio
async def test_a_broken_file_is_marked_failed() -> None:
    service, _, _ = _service(docling=None)
    attachment = await _upload(service, b"not a zip archive", "memo.docx")
    assert attachment["extraction_status"] == ExtractionStatus.FAILED


@pytest.mark.anyio
async def test_a_type_nothing_can_read_is_refused_before_it_is_stored(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "chat_upload_allowed_extensions", ".odt")
    service, dao, store = _service(docling=None)
    with pytest.raises(AttachmentError):
        await _upload(service, b"...", "letter.odt")
    assert store.keys == []
    assert dao.created == []


def test_extract_text_refuses_unknown_extensions() -> None:
    with pytest.raises(ValueError, match="odt"):
        extraction.extract_text(b"...", "letter.odt")
