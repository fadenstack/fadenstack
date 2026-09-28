"""Local text extraction for chat attachments.

Used when the Docling service is not configured or fails, so a PDF or an
Office file attached to a chat still reaches the model as text. Chat
attachments needed Docling for anything but plain text; without it the
upload was refused, and with Docling down the attachment was stored with
no text at all.

The same small libraries as the backend's fallback for knowledge
ingestion (``faden_backend/services/docling/fallback.py``):

* **PDF** -> ``pdfplumber``: text and simple tables, no OCR.
* **DOCX** -> ``python-docx``: paragraphs and tables.
* **PPTX** -> ``python-pptx``: text frames per slide.
* **XLSX** -> ``openpyxl``: cell values per sheet.
* **CSV** -> stdlib ``csv``.

No OCR, no layout: a scanned PDF yields no text. Docling stays the better
choice where it runs.
"""

from __future__ import annotations

import csv
import io
import logging
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

#: Extensions this module can extract.
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx", ".xlsx", ".csv"})


@dataclass
class Extracted:
    """Text from one file, in the shape the Docling client returns."""

    content: str
    #: Pages (PDF) or slides (PPTX) in the file, not only those read.
    page_count: int | None = None
    format_detected: str = "unknown"
    tables_found: int = 0
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        """``{content, metadata}``, as ``ChatDoclingClient.convert`` returns it."""
        return {
            "content": self.content,
            "metadata": {
                "page_count": self.page_count,
                "format_detected": self.format_detected,
                "tables_found": self.tables_found,
                "ocr_applied": False,
                "backend": "local",
                **self.metadata,
            },
        }


def _pdf(data: bytes, max_pages: int | None) -> Extracted:
    import pdfplumber  # noqa: PLC0415

    pages: list[str] = []
    tables_found = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        total = len(pdf.pages)
        for page in pdf.pages[:max_pages]:
            parts: list[str] = []
            tables = page.find_tables()
            tables_found += len(tables)
            for table in tables:
                rows = [
                    " | ".join(str(cell) if cell else "" for cell in row)
                    for row in table.extract()
                    if any(row)
                ]
                if rows:
                    parts.append("\n".join(rows))
            # Text outside the tables, so table cells are not read twice.
            rest = page
            for table in tables:
                rest = rest.outside_bbox(table.bbox)
            text = (rest.extract_text() or "").strip()
            if text:
                parts.insert(0, text)
            if parts:
                pages.append("\n\n".join(parts))
    return Extracted(
        "\n\n".join(pages),
        page_count=total,
        format_detected="pdf",
        tables_found=tables_found,
    )


def _docx(data: bytes, max_pages: int | None) -> Extracted:
    from docx import Document  # noqa: PLC0415

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        parts.append(
            "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ),
        )
    return Extracted(
        "\n\n".join(parts),
        format_detected="docx",
        tables_found=len(doc.tables),
    )


def _pptx(data: bytes, max_pages: int | None) -> Extracted:
    from pptx import Presentation  # noqa: PLC0415

    prs = Presentation(io.BytesIO(data))
    slides_all = list(prs.slides)
    slides: list[str] = []
    for slide in slides_all[:max_pages]:
        lines = [
            para.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            if para.text.strip()
        ]
        if lines:
            slides.append("\n".join(lines))
    return Extracted(
        "\n\n---\n\n".join(slides),
        page_count=len(slides_all),
        format_detected="pptx",
    )


def _xlsx(data: bytes, max_pages: int | None) -> Extracted:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"## {ws.title}\n\n" + "\n".join(rows))
    finally:
        wb.close()
    return Extracted(
        "\n\n".join(sheets),
        format_detected="xlsx",
        tables_found=len(sheets),
    )


def _csv(data: bytes, max_pages: int | None) -> Extracted:
    reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
    rows = [" | ".join(row) for row in reader if any(row)]
    return Extracted("\n".join(rows), format_detected="csv", tables_found=int(bool(rows)))


_EXTRACTORS: dict[str, Callable[[bytes, int | None], Extracted]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".xlsx": _xlsx,
    ".csv": _csv,
}


def extract_text(
    file_bytes: bytes,
    filename: str,
    *,
    max_pages: int | None = None,
) -> Extracted:
    """Extract the text of *file_bytes*, reading at most *max_pages* pages.

    Raises ``ValueError`` for an extension this module does not handle, and
    lets a library's error for a broken file propagate: the caller decides
    what a failed extraction means.
    """
    ext = Path(filename).suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(f"No local extractor for '{ext}' files.")
    return extractor(file_bytes, max_pages)
